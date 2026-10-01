"""Apuntes a Word: convierte PDFs de apuntes en documentos Word de estudio.

Por cada PDF (o por todos juntos) genera un .docx con:
  - Índice
  - Apuntes pasados a limpio, con las imágenes del PDF en su sitio
  - Resumen
  - Esquema
  - Preguntas tipo test (con soluciones explicadas al final)
  - Fichas de repaso (pregunta / respuesta)

Una segunda pasada de revisión compara el texto con el PDF y completa lo que falte.
"""

import base64
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import threading
import time
import traceback
import tkinter as tk
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

import anthropic
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader, PdfWriter

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except Exception:  # sin arrastrar y soltar si la librería no está
    TkinterDnD = None

MODELO = "claude-opus-5-5"
ESFUERZO = "high"
PRECIO = {"entrada": 4.0, "salida": 20.0, "cache_lectura": 0.20, "cache_escritura": 5.0}  # US$ por millón de tokens
PAGINAS_POR_TROZO = 40          # trozos pequeños: salidas más completas y fieles
MAX_BYTES_TROZO = 24 * 1024 * 1024
AREA_MIN_IMAGEN = 40_000        # px²; por debajo son iconos
if sys.platform == "darwin":
    CONFIG_DIR = Path.home() / "Library" / "Application Support" / "ApuntesAWord"
else:
    CONFIG_DIR = Path(os.environ.get("APPDATA", Path.home())) / "ApuntesAWord"
CONFIG_FILE = CONFIG_DIR / "config.json"

SISTEMA = (
    "Eres un asistente experto que ayuda a una estudiante universitaria a preparar sus apuntes. "
    "Escribe siempre en español de España, con ortografía y tildes perfectas. "
    "Sé fiel al material: no inventes datos que no estén en los apuntes. "
    "Responde solo con Markdown (títulos con ##, listas con -, tablas con |, negritas con **), "
    "sin frases introductorias, comentarios ni despedidas."
)

INSTRUCCIONES_LIMPIO = """Pásalos a limpio como un documento de estudio completo, bien redactado y ordenado:
- Conserva TODO el contenido y todos los datos (cifras, valores normales, clasificaciones, definiciones, causas, consecuencias, cuidados de enfermería, bibliografía). Ninguna diapositiva puede quedar sin reflejar.
- Mantén el orden del PDF.
- Une frases cortadas por el formato de diapositiva y elimina títulos repetidos en cada página, pies de página y logotipos.
- No pongas un título general. Estructura:
  - Apartados principales numerados y en MAYÚSCULAS: "## {inicio}. TÍTULO DEL APARTADO", "## {siguiente}. OTRO APARTADO"...
  - Subapartados en MAYÚSCULAS sin número: "### SUBAPARTADO".
  - Títulos menores (nombres de índices, fórmulas, clasificaciones): "#### Título".
- Estilo de redacción: frases breves en listas con viñetas (hasta 3 niveles con 2 espacios de sangría por nivel). Cuando un elemento tenga un término y su explicación, escribe el término en negrita con dos puntos: "- **Alta vulnerabilidad:** los niños requieren...". Usa párrafos solo para introducciones cortas.
- Tablas cuando la información sea tabular (comparaciones, tipos con descripción y riesgos...). No dejes celdas vacías.
- Recuadros: pon en líneas que empiecen por "> " los datos que conviene memorizar de un vistazo (fórmulas, valores normales, constantes, cifras clave, definiciones esenciales). Una idea por línea.
- Lo MUY importante (lo que seguro cae en el examen) enciérralo entre == ==, por ejemplo ==MUY IMPORTANTE: ...==. Úsalo con moderación (pocas veces por tema).
- Resalta en **negrita** los conceptos clave.
- Si una imagen contiene información que hay que estudiar (rótulos, datos, esquemas), transcríbela en texto además de colocar su marcador.
{imagenes}"""

INSTRUCCIONES_IMAGENES = """- Imágenes disponibles (identificador, página, tamaño en píxeles):
{lista}
  Donde una imagen aporte valor para estudiar, colócala en su propia línea con el formato [[IMG_página_número|pie de foto breve]], por ejemplo [[IMG_13_1|Placenta con el cordón umbilical]], justo después del texto al que ilustra. Omite las decorativas (portadas, fondos, logotipos)."""

PROMPT_REVISION = """Arriba está el PDF original. Debajo tienes un borrador de estos apuntes pasados a limpio.
Revísalo página por página contra el PDF, como un corrector exigente:
- Añade todo lo que falte (datos, apartados, filas de tablas, elementos de listas, información de las imágenes).
- Corrige cualquier error, dato cambiado o frase incompleta, y rellena las celdas de tabla vacías.
- Comprueba que no quede ningún apartado vacío ni frase que termine en dos puntos sin nada detrás.
- Conserva los marcadores [[IMG_...|...]] y añade los de imágenes útiles que falten.
- Mantén el mismo formato: apartados "## N. TÍTULO", subapartados "### ...", recuadros "> ", importante == ==.
Devuelve el documento completo ya corregido (no una lista de cambios).
{imagenes}
<borrador>
{borrador}
</borrador>"""

SECCIONES_ESTUDIO = {
    "resumen": (
        "## Resumen\n"
        "Un resumen claro de todos los apuntes (aprox. una cuarta parte de su extensión), con "
        "'### N. BLOQUE EN MAYÚSCULAS' para cada bloque temático, listas breves, los conceptos "
        "clave en **negrita** y los datos para memorizar en líneas que empiecen por '> '."
    ),
    "esquema": (
        "## Esquema\n"
        "Un esquema jerárquico de todo el temario con listas anidadas (usa - y 2 espacios de "
        "sangría por nivel, máximo 3 niveles), con palabras clave, no frases largas."
    ),
    "test": (
        "## Preguntas tipo test\n"
        "Entre 20 y 30 preguntas tipo test numeradas (1., 2., ...) que cubran todo el temario, "
        "cada una con 4 opciones en líneas separadas que empiecen por a), b), c), d). Solo una "
        "correcta, respaldada por los apuntes, con distractores verosímiles. Al final, un apartado "
        "'### Soluciones' con una línea por pregunta: '1. b) explicación breve'."
    ),
    "fichas": (
        "## Fichas de repaso\n"
        "Entre 20 y 35 fichas. Cada ficha en dos líneas: '**P:** pregunta' y '**R:** respuesta "
        "breve', separadas por una línea en blanco."
    ),
}
TITULOS_ESTUDIO = ("Resumen", "Esquema", "Preguntas tipo test", "Fichas de repaso")


# ----------------------------------------------------------------------------- configuración

def carpeta_descargas():
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r"Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders") as k:
            return winreg.QueryValueEx(k, "{374DE290-123F-4565-9164-39C4925E467B}")[0]
    except Exception:
        return str(Path.home() / "Downloads")


def limpiar_clave(texto):
    """Saca la clave aunque se pegue con más texto alrededor (p. ej. el ejemplo curl de la web)."""
    m = re.search(r"sk-ant-[A-Za-z0-9_\-]+", texto or "")
    return m.group(0) if m else (texto or "").strip()


def cargar_config():
    try:
        cfg = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if cfg.get("api_key"):
        cfg["api_key"] = limpiar_clave(cfg["api_key"])
    return cfg


def guardar_config(cfg):
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2), encoding="utf-8")


# ----------------------------------------------------------------------------- PDF

def extraer_imagenes(lector):
    """{id: PIL.Image} con las imágenes útiles (sin logos repetidos ni iconos)."""
    encontradas = []
    repeticiones = Counter()
    for num, pagina in enumerate(lector.pages, 1):
        try:
            imagenes = list(pagina.images)
        except Exception:
            continue
        for im in imagenes:
            try:
                h = hashlib.md5(im.data).hexdigest()
                repeticiones[h] += 1
                encontradas.append((num, h, im))
            except Exception:
                continue
    resultado, vistas, por_pagina = {}, set(), Counter()
    for num, h, im in encontradas:
        if repeticiones[h] > 3 or h in vistas:
            continue
        try:
            pil = im.image
            if pil.width * pil.height < AREA_MIN_IMAGEN:
                continue
        except Exception:
            continue
        vistas.add(h)
        por_pagina[num] += 1
        resultado[f"IMG_{num}_{por_pagina[num]}"] = pil
    return resultado


def trocear_pdf(datos, lector):
    """Lista de (pdf_bytes, página_inicio, página_fin)."""
    total = len(lector.pages)
    if total <= PAGINAS_POR_TROZO and len(datos) <= MAX_BYTES_TROZO:
        return [(datos, 1, total)]
    trozos = []
    for inicio in range(0, total, PAGINAS_POR_TROZO):
        fin = min(inicio + PAGINAS_POR_TROZO, total)
        escritor = PdfWriter()
        for i in range(inicio, fin):
            escritor.add_page(lector.pages[i])
        buf = io.BytesIO()
        escritor.write(buf)
        trozos.append((buf.getvalue(), inicio + 1, fin))
    return trozos


def texto_imagenes(imagenes, ini, fin):
    lineas = []
    for ident, im in imagenes.items():
        pag = int(ident.split("_")[1])
        if ini <= pag <= fin:
            lineas.append(f"  {ident} (pág. {pag}, {im.width}x{im.height})")
    if not lineas:
        return ""
    return INSTRUCCIONES_IMAGENES.format(lista="\n".join(lineas))


# ----------------------------------------------------------------------------- llamadas a Claude

class Contador:
    def __init__(self):
        self.coste = 0.0
        self.cerrojo = threading.Lock()

    def sumar(self, u):
        m = 1_000_000
        with self.cerrojo:
            self.coste += (
            (u.input_tokens or 0) * PRECIO["entrada"] / m
            + (u.output_tokens or 0) * PRECIO["salida"] / m
            + (getattr(u, "cache_read_input_tokens", 0) or 0) * PRECIO["cache_lectura"] / m
            + (getattr(u, "cache_creation_input_tokens", 0) or 0) * PRECIO["cache_escritura"] / m
        )


def pedir(cliente, contenido, contador):
    """Una petición en streaming (salidas largas) con respaldo automático si el modelo rechaza."""
    with cliente.beta.messages.stream(
        model=MODELO,
        max_tokens=64000,
        system=SISTEMA,
        output_config={"effort": ESFUERZO},
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        messages=[{"role": "user", "content": contenido}],
    ) as stream:
        respuesta = stream.get_final_message()
    contador.sumar(respuesta.usage)
    if respuesta.stop_reason == "refusal":
        raise RuntimeError("La IA no ha podido procesar este documento.")
    texto = "".join(b.text for b in respuesta.content if b.type == "text").strip()
    texto = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", texto)
    if respuesta.stop_reason == "max_tokens":
        texto += "\n\n(El texto se ha cortado por ser demasiado largo.)"
    return texto


def bloque_pdf(pdf):
    return {
        "type": "document",
        "source": {"type": "base64", "media_type": "application/pdf",
                   "data": base64.standard_b64encode(pdf).decode("ascii")},
        "cache_control": {"type": "ephemeral"},   # la revisión reutiliza el PDF más barato
    }


def renumerar(md, nivel="##"):
    """Numera los apartados seguidos (1., 2., 3....) aunque vengan de trozos distintos."""
    contador = iter(range(1, 10_000))
    return re.sub(rf"(?m)^{nivel}\s+(?:\d+\s*[.)]\s*)?(.+)$",
                  lambda m: f"{nivel} {next(contador)}. {m.group(1)}", md)


def pasar_a_limpio(cliente, datos, lector, imagenes, revisar, contador, log):
    trozos = trocear_pdf(datos, lector)

    def un_trozo(trozo):
        pdf, ini, fin = trozo
        etiqueta = f" (páginas {ini}-{fin})" if len(trozos) > 1 else ""
        extra_img = texto_imagenes(imagenes, ini, fin)
        contexto = (f"Este PDF es el fragmento de las páginas {ini} a {fin} de los apuntes; "
                    f"su primera página es la página {ini}.\n" if len(trozos) > 1 else "")
        log(f"   Pasando a limpio{etiqueta}...")
        prompt = (contexto + "Estos son apuntes en PDF (probablemente diapositivas de clase). "
                  + INSTRUCCIONES_LIMPIO.format(imagenes=extra_img, inicio=1, siguiente=2))
        borrador = pedir(cliente, [bloque_pdf(pdf), {"type": "text", "text": prompt}], contador)
        if revisar:
            log(f"   Revisando contra el PDF{etiqueta}...")
            prompt_rev = contexto + PROMPT_REVISION.format(imagenes=extra_img, borrador=borrador)
            borrador = pedir(cliente, [bloque_pdf(pdf), {"type": "text", "text": prompt_rev}], contador)
        log(f"   ✓ Terminado{etiqueta}")
        return borrador

    # los trozos se procesan a la vez para tardar menos
    with ThreadPoolExecutor(max_workers=4) as grupo:
        partes = list(grupo.map(un_trozo, trozos))
    return renumerar("\n\n".join(partes))


def material_estudio(cliente, apuntes_md, opciones, contador):
    pedidas = [SECCIONES_ESTUDIO[k] for k in ("resumen", "esquema", "test", "fichas") if opciones[k]]
    if not pedidas:
        return ""
    apuntes_sin_img = re.sub(r"\[\[IMG_[^\]]*\]\]", "", apuntes_md)

    def una_seccion(instruccion):
        prompt = (
            "<apuntes>\n" + apuntes_sin_img + "\n</apuntes>\n\n"
            "A partir de estos apuntes, crea este material de estudio. Empieza exactamente con el "
            "título indicado y no añadas otros apartados:\n\n" + instruccion
            + "\n\nComprueba antes de responder que todo está respaldado por los apuntes y que "
              "ningún dato los contradice."
        )
        return pedir(cliente, [{"type": "text", "text": prompt}], contador)

    # cada apartado se genera a la vez para tardar menos
    with ThreadPoolExecutor(max_workers=4) as grupo:
        return "\n\n".join(grupo.map(una_seccion, pedidas))


# ----------------------------------------------------------------------------- Markdown -> Word

# Estilo copiado de los apuntes de ella (ENFERMERÍA PEDIÁTRICA.docx)
FUENTE_TEXTO = "Aptos"
FUENTE_TITULOS = "Modern Love Caps"
LILA = "F2CFEE"
TAM_TEXTO = Pt(10.5)

INLINE = re.compile(r"(\*\*.+?\*\*|==.+?==|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")
MARCA_IMG = re.compile(r"\[\[(IMG_\d+_\d+)(?:\|([^\]]*))?\]\]")


def añadir_texto(parrafo, texto):
    texto = texto.replace("<br>", " ").replace("<br/>", " ")
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            interior = trozo[2:-2]
            r = parrafo.add_run(interior.replace("==", ""))
            r.bold = True
            if interior.startswith("==") and interior.endswith("=="):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif trozo.startswith("==") and trozo.endswith("==") and len(trozo) > 4:
            r = parrafo.add_run(trozo[2:-2].replace("**", ""))
            r.bold = True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo.replace("**", ""))


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _borde(elemento_padre, etiqueta, lados, color, grosor=4):
    bdr = OxmlElement(etiqueta)
    for lado in lados:
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(grosor))
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), color)
        bdr.append(b)
    elemento_padre.append(bdr)


def caja_parrafo(pPr, relleno=None, borde="000000"):
    """Recuadro alrededor del párrafo (como las cajas de texto de sus apuntes)."""
    _borde(pPr, "w:pBdr", ("top", "left", "bottom", "right"), borde)
    if relleno:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), relleno)
        pPr.append(shd)


def fuente(estilo, nombre):
    estilo.font.name = nombre
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for atr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atr), nombre)


def viñetas_como_word(doc):
    """Niveles de viñeta • / o / ▪ como en Word."""
    numbering = doc.part.numbering_part.element
    simbolos = {"List Bullet": ("", "Symbol"), "List Bullet 2": ("o", "Courier New"),
                "List Bullet 3": ("", "Wingdings")}
    for estilo, (texto, letra) in simbolos.items():
        num_id = doc.styles[estilo].element.xpath(".//w:numId/@w:val")
        if not num_id:
            continue
        abs_id = numbering.xpath(f'./w:num[@w:numId="{num_id[0]}"]/w:abstractNumId/@w:val')
        if not abs_id:
            continue
        for lvl in numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abs_id[0]}"]/w:lvl[@w:ilvl="0"]'):
            lvl.find(qn("w:lvlText")).set(qn("w:val"), texto)
            for rf in lvl.iter(qn("w:rFonts")):
                for atr in ("w:ascii", "w:hAnsi", "w:cs"):
                    rf.set(qn(atr), letra)


def aplicar_estilo(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    fuente(normal, FUENTE_TEXTO)
    normal.font.size = TAM_TEXTO
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    pf.line_spacing = 1.16

    # Título del documento
    t = doc.styles["Title"]
    fuente(t, FUENTE_TITULOS)
    t.font.size = Pt(22)
    t.font.color.rgb = RGBColor(0, 0, 0)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Nivel 1: bloques del documento (Apuntes, Resumen, Test...)
    h1 = doc.styles["Heading 1"]
    fuente(h1, FUENTE_TITULOS)
    h1.font.size = Pt(16)
    h1.font.bold = False
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_after = Pt(12)
    _borde(h1.element.get_or_add_pPr(), "w:pBdr", ("bottom",), "A02B93", 8)

    # Nivel 2: "1. CONCEPTOS" en caja lila, como sus cuadros de texto
    h2 = doc.styles["Heading 2"]
    fuente(h2, FUENTE_TITULOS)
    h2.font.size = Pt(11)
    h2.font.bold = False
    h2.font.all_caps = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    caja_parrafo(h2.element.get_or_add_pPr(), relleno=LILA, borde=LILA)

    # Nivel 3: subapartado en su fuente decorativa, subrayado
    h3 = doc.styles["Heading 3"]
    fuente(h3, FUENTE_TITULOS)
    h3.font.size = TAM_TEXTO
    h3.font.bold = False
    h3.font.underline = True
    h3.font.all_caps = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(10)

    # Nivel 4: negrita subrayada (como "TASA BRUTA DE NATALIDAD")
    h4 = doc.styles["Heading 4"]
    fuente(h4, FUENTE_TEXTO)
    h4.font.size = TAM_TEXTO
    h4.font.bold = True
    h4.font.italic = False
    h4.font.underline = True
    h4.font.color.rgb = RGBColor(0, 0, 0)
    h4.paragraph_format.space_before = Pt(8)

    for nombre in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number"):
        est = doc.styles[nombre]
        est.paragraph_format.space_after = Pt(2)
        est.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    viñetas_como_word(doc)


def añadir_tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    celdas = [f for f in celdas if not all(re.fullmatch(r":?-{2,}:?", c) for c in f if c)]
    if not celdas:
        return
    ncols = max(len(f) for f in celdas)
    tabla = doc.add_table(rows=len(celdas), cols=ncols)
    tabla.style = "Table Grid"
    # bordes grises suaves, como su "Tabla de cuadrícula 2"
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(viejo)
    _borde(tblPr, "w:tblBorders", ("top", "left", "bottom", "right", "insideH", "insideV"), "BFBFBF")
    for i, fila in enumerate(celdas):
        for j in range(ncols):
            celda = tabla.cell(i, j)
            p = celda.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            añadir_texto(p, fila[j] if j < len(fila) else "")
            for r in p.runs:
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True
            if i == 0:
                sombrear(celda, LILA)
            elif i % 2 == 0:
                sombrear(celda, "F2F2F2")
    doc.add_paragraph()


def añadir_imagen(doc, pil, pie):
    if pil.mode not in ("RGB", "L"):
        fondo = Image.new("RGB", pil.size, "white")
        rgba = pil.convert("RGBA")
        fondo.paste(rgba, mask=rgba.split()[-1])
        pil = fondo
    if max(pil.size) > 1400:
        pil = pil.copy()
        pil.thumbnail((1400, 1400))
    buf = io.BytesIO()
    pil.convert("RGB").save(buf, format="JPEG", quality=85)
    buf.seek(0)
    ancho = min(14.0, max(5.0, pil.width / 45))
    alto_max = 11.0
    if ancho * pil.height / pil.width > alto_max:
        ancho = alto_max * pil.width / pil.height
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(buf, width=Cm(ancho))
    if pie:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(pie.strip())
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def markdown_a_docx(doc, md, imagenes=None, desplazamiento=0):
    imagenes = imagenes or {}
    usadas = set()
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        s = linea.strip()
        if not s:
            i += 1
            continue
        m_img = MARCA_IMG.search(s)
        if m_img:
            ident, pie = m_img.group(1), m_img.group(2)
            if ident in imagenes and ident not in usadas:
                usadas.add(ident)
                try:
                    añadir_imagen(doc, imagenes[ident], pie)
                except Exception:
                    pass
            s = MARCA_IMG.sub("", s).strip()
            linea = MARCA_IMG.sub("", linea)
            if not s:
                i += 1
                continue
        if s.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i])
                i += 1
            añadir_tabla(doc, filas)
            continue
        if s.startswith(">"):
            # líneas "> " seguidas = un único recuadro con borde negro fino
            p = None
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                contenido = lineas[i].strip().lstrip(">").strip()
                if contenido:
                    if p is None:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_before = Pt(6)
                        caja_parrafo(p._p.get_or_add_pPr())
                    else:
                        p.add_run().add_break()
                    añadir_texto(p, re.sub(r"^[-*•]\s+", "", contenido))
                i += 1
            if p is not None:
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            nivel = max(2, min(len(m.group(1)) + desplazamiento, 4))
            h = doc.add_heading(level=nivel)
            añadir_texto(h, m.group(2).strip("# "))
        elif re.fullmatch(r"[-*_]{3,}", s):
            pass
        else:
            sangria = (len(linea) - len(linea.lstrip(" "))) // 2
            mb = re.match(r"^[-*•▪]\s+(.*)", s)
            mn = re.match(r"^(\d+)[.)]\s+(.*)", s)
            mo = re.match(r"^([a-dA-D])\)\s+(.*)", s)
            if mb:
                estilo = "List Bullet" if sangria == 0 else f"List Bullet {min(sangria + 1, 3)}"
                añadir_texto(doc.add_paragraph(style=estilo), mb.group(1))
            elif mn:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.add_run(mn.group(1) + ". ").bold = True
                añadir_texto(p, mn.group(2))
            elif mo:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(0)
                añadir_texto(p, f"{mo.group(1)}) {mo.group(2)}")
            else:
                añadir_texto(doc.add_paragraph(), s)
        i += 1


def salto_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def numeros_de_pagina(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("pág. ")
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def añadir_indice(doc, partes, estudio_md):
    doc.add_heading("Índice", level=1)
    for nombre, md in partes:
        if nombre:
            doc.add_paragraph(nombre, style="List Number").runs[0].bold = True
        for t in re.findall(r"(?m)^##\s+(.+)$", md):
            doc.add_paragraph(t.strip("# ").replace("**", ""),
                              style="List Bullet 2" if nombre else "List Bullet")
    for t in TITULOS_ESTUDIO:
        if re.search(rf"(?m)^##\s+{t}", estudio_md or ""):
            p = doc.add_paragraph(t, style="List Number" if partes[0][0] else "List Bullet")
            p.runs[0].bold = True


def crear_word(titulo, partes, estudio_md, imagenes, destino):
    """partes: lista de (subtítulo o None, markdown de apuntes)."""
    doc = Document()
    aplicar_estilo(doc)
    numeros_de_pagina(doc)

    doc.add_heading(titulo, level=0)
    añadir_indice(doc, partes, estudio_md)

    for nombre, md in partes:
        if not md:
            continue
        salto_pagina(doc)
        doc.add_heading(nombre or "Apuntes", level=1)
        markdown_a_docx(doc, md, imagenes)
    if estudio_md:
        patron = r"(?m)^##\s+(?=" + "|".join(TITULOS_ESTUDIO) + ")"
        for bloque in re.split(patron, estudio_md):
            cabecera, _, resto = bloque.partition("\n")
            if not cabecera.strip().startswith(TITULOS_ESTUDIO):
                continue
            salto_pagina(doc)
            doc.add_heading(cabecera.strip(), level=1)
            markdown_a_docx(doc, resto, desplazamiento=-1)
    doc.save(destino)


# ----------------------------------------------------------------------------- proceso completo

def ruta_libre(carpeta, nombre):
    destino = Path(carpeta) / f"{nombre}.docx"
    n = 2
    while destino.exists():
        destino = Path(carpeta) / f"{nombre} ({n}).docx"
        n += 1
    return destino


def leer_pdf(ruta):
    datos = Path(ruta).read_bytes()
    lector = PdfReader(io.BytesIO(datos))
    if lector.is_encrypted:
        lector.decrypt("")
    return datos, lector


def procesar(rutas, carpeta, opciones, api_key, log, unir=False):
    """Procesa los PDFs. Devuelve (lista de Word creados, coste en US$)."""
    cliente = anthropic.Anthropic(api_key=api_key, max_retries=5)
    contador = Contador()
    grupos = [rutas] if unir and len(rutas) > 1 else [[r] for r in rutas]
    creados = []
    for grupo in grupos:
        partes, imagenes_todas = [], {}
        for ruta in grupo:
            nombre = Path(ruta).stem
            inicio = time.time()
            log(f"▶ {nombre}")
            datos, lector = leer_pdf(ruta)
            imagenes = extraer_imagenes(lector) if opciones["imagenes"] else {}
            log(f"   {len(lector.pages)} páginas, {len(imagenes)} imágenes útiles")
            md =pasar_a_limpio(cliente, datos, lector, imagenes, opciones["revisar"], contador, log)
            if len(grupo) > 1:
                sufijo = f"_{len(partes) + 1}"
                md = MARCA_IMG.sub(lambda m: f"[[{m.group(1)}{sufijo}|{m.group(2) or ''}]]", md)
                imagenes = {k + sufijo: v for k, v in imagenes.items()}
            imagenes_todas.update(imagenes)
            partes.append((nombre if len(grupo) > 1 else None, md))
            log(f"   Hecho en {int(time.time() - inicio)} s")
        estudio = ""
        if any(opciones[k] for k in ("resumen", "esquema", "test", "fichas")):
            log("   Creando resumen, esquema, test y fichas...")
            estudio = material_estudio(cliente, "\n\n".join(md for _, md in partes), opciones, contador)
        if not opciones["limpio"]:
            partes = [(n, "") for n, _ in partes]
        titulo = Path(grupo[0]).stem if len(grupo) == 1 else f"Apuntes: {Path(grupo[0]).stem} y {len(grupo) - 1} más"
        destino = ruta_libre(carpeta, titulo if len(grupo) == 1 else "Apuntes unidos")
        crear_word(titulo, partes, estudio, imagenes_todas, destino)
        log(f"   ✔ Guardado: {destino}")
        creados.append(destino)
    return creados, contador.coste


def abrir(ruta):
    if sys.platform == "darwin":
        if subprocess.call(["open", str(ruta)]) != 0:  # sin programa para .docx: se muestra en Finder
            subprocess.Popen(["open", "-R", str(ruta)])
        return
    try:
        os.startfile(ruta)
    except OSError:  # no hay programa para .docx: se muestra en la carpeta
        subprocess.Popen(["explorer", "/select,", str(ruta)])


# ----------------------------------------------------------------------------- interfaz

class App:
    OPCIONES = [
        ("limpio", "Apuntes pasados a limpio"),
        ("imagenes", "Incluir imágenes del PDF"),
        ("revisar", "Revisión extra contra el PDF"),
        ("resumen", "Resumen"),
        ("esquema", "Esquema"),
        ("test", "Preguntas tipo test"),
        ("fichas", "Fichas de repaso"),
    ]

    def __init__(self, raiz, pdfs_iniciales=()):
        self.raiz = raiz
        self.cfg = cargar_config()
        self.pdfs = []
        raiz.title("Apuntes a Word")
        raiz.geometry("680x640")
        raiz.minsize(600, 560)

        marco = ttk.Frame(raiz, padding=14)
        marco.pack(fill="both", expand=True)

        ttk.Label(marco, text="Apuntes a Word", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(marco, text="Elige o arrastra tus PDFs y te preparo el Word para estudiar.").pack(anchor="w", pady=(0, 8))

        caja = ttk.LabelFrame(marco, text="PDFs", padding=8)
        caja.pack(fill="x")
        self.lista = tk.Listbox(caja, height=4, activestyle="none")
        self.lista.pack(fill="x")
        self.lista.insert("end", "  Arrastra aquí tus PDFs o pulsa «Añadir PDFs»")
        botones = ttk.Frame(caja)
        botones.pack(fill="x", pady=(6, 0))
        ttk.Button(botones, text="📄 Añadir PDFs...", command=self.elegir_pdfs).pack(side="left")
        ttk.Button(botones, text="Quitar todos", command=self.vaciar).pack(side="left", padx=6)
        self.unir = tk.BooleanVar(value=self.cfg.get("unir", False))
        ttk.Checkbutton(botones, text="Unir todos en un solo Word", variable=self.unir).pack(side="right")

        opciones = ttk.LabelFrame(marco, text="¿Qué quieres en el Word?", padding=8)
        opciones.pack(fill="x", pady=10)
        self.vars = {}
        guardadas = self.cfg.get("opciones", {})
        for n, (clave, texto) in enumerate(self.OPCIONES):
            v = tk.BooleanVar(value=guardadas.get(clave, True))
            self.vars[clave] = v
            ttk.Checkbutton(opciones, text=texto, variable=v).grid(row=n // 3, column=n % 3, sticky="w", padx=6, pady=2)

        fila2 = ttk.Frame(marco)
        fila2.pack(fill="x")
        ttk.Label(fila2, text="Guardar en:").pack(side="left")
        self.carpeta = tk.StringVar(value=self.cfg.get("carpeta_salida") or carpeta_descargas())
        ttk.Entry(fila2, textvariable=self.carpeta).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(fila2, text="Cambiar...", command=self.elegir_carpeta).pack(side="left")

        fila3 = ttk.Frame(marco)
        fila3.pack(fill="x", pady=10)
        self.btn = ttk.Button(fila3, text="✨ Crear Word", command=self.empezar)
        self.btn.pack(side="left", ipadx=10, ipady=3)
        self.abrir_al_final = tk.BooleanVar(value=self.cfg.get("abrir", True))
        ttk.Checkbutton(fila3, text="Abrir el Word al terminar", variable=self.abrir_al_final).pack(side="left", padx=12)
        ttk.Button(fila3, text="⚙ Clave de la IA", command=self.pedir_clave).pack(side="right")

        self.progreso = ttk.Progressbar(marco, mode="indeterminate")
        self.progreso.pack(fill="x")
        self.registro = scrolledtext.ScrolledText(marco, height=10, state="disabled", font=("Segoe UI", 9))
        self.registro.pack(fill="both", expand=True, pady=(8, 0))

        if TkinterDnD is not None:
            for w in (raiz, self.lista):
                w.drop_target_register(DND_FILES)
                w.dnd_bind("<<Drop>>", self.soltar)

        self.añadir(pdfs_iniciales)
        if not self.cfg.get("api_key"):
            raiz.after(300, self.pedir_clave)

    # ---- lista de PDFs
    def añadir(self, rutas):
        nuevas = [r for r in rutas if r.lower().endswith(".pdf") and r not in self.pdfs]
        if not nuevas:
            return
        if not self.pdfs:
            self.lista.delete(0, "end")
        for r in nuevas:
            self.pdfs.append(r)
            self.lista.insert("end", "  " + Path(r).name)

    def vaciar(self):
        self.pdfs = []
        self.lista.delete(0, "end")
        self.lista.insert("end", "  Arrastra aquí tus PDFs o pulsa «Añadir PDFs»")

    def soltar(self, evento):
        self.añadir(list(self.raiz.tk.splitlist(evento.data)))

    def elegir_pdfs(self):
        self.añadir(list(filedialog.askopenfilenames(title="Elige tus apuntes", filetypes=[("PDF", "*.pdf")])))

    def elegir_carpeta(self):
        c = filedialog.askdirectory(title="¿Dónde guardo los Word?", initialdir=self.carpeta.get())
        if c:
            self.carpeta.set(c)

    # ---- registro
    def log(self, texto):
        def escribir():
            self.registro.configure(state="normal")
            self.registro.insert("end", texto + "\n")
            self.registro.see("end")
            self.registro.configure(state="disabled")
        self.raiz.after(0, escribir)

    # ---- clave
    def pedir_clave(self):
        ventana = tk.Toplevel(self.raiz)
        ventana.title("Clave de la IA")
        ventana.transient(self.raiz)
        ventana.grab_set()
        f = ttk.Frame(ventana, padding=14)
        f.pack(fill="both", expand=True)
        ttk.Label(f, text="Pega aquí la clave de Anthropic (empieza por sk-ant-):").pack(anchor="w")
        v = tk.StringVar(value=self.cfg.get("api_key", ""))
        e = ttk.Entry(f, textvariable=v, width=60, show="•")
        e.pack(fill="x", pady=8)
        e.focus()

        def guardar():
            self.cfg["api_key"] = limpiar_clave(v.get())
            guardar_config(self.cfg)
            ventana.destroy()

        ttk.Button(f, text="Guardar", command=guardar).pack(anchor="e")

    # ---- proceso
    def empezar(self):
        if not self.pdfs:
            messagebox.showinfo("Apuntes a Word", "Primero añade al menos un PDF.")
            return
        if not self.cfg.get("api_key"):
            self.pedir_clave()
            return
        opciones = {k: v.get() for k, v in self.vars.items()}
        if not any(opciones[k] for k in ("limpio", "resumen", "esquema", "test", "fichas")):
            messagebox.showinfo("Apuntes a Word", "Marca al menos una opción.")
            return
        carpeta = self.carpeta.get()
        try:
            Path(carpeta).mkdir(parents=True, exist_ok=True)
        except OSError:
            messagebox.showerror("Apuntes a Word", "No puedo usar esa carpeta. Elige otra.")
            return
        self.cfg.update(opciones=opciones, carpeta_salida=carpeta, unir=self.unir.get(),
                        abrir=self.abrir_al_final.get())
        guardar_config(self.cfg)

        self.btn.config(state="disabled")
        self.progreso.start(12)
        minutos = 5 if opciones["revisar"] else 3
        self.log(f"Empiezo. Cuenta unos {minutos} minutos por PDF; puedes dejarlo trabajando.")
        threading.Thread(target=self.trabajar, args=(list(self.pdfs), carpeta, opciones, self.unir.get()),
                         daemon=True).start()

    def trabajar(self, pdfs, carpeta, opciones, unir):
        creados, coste = [], 0.0
        try:
            creados, coste = procesar(pdfs, carpeta, opciones, self.cfg["api_key"], self.log, unir)
        except anthropic.AuthenticationError:
            self.log("✖ La clave de la IA no es válida. Revísala en ⚙ Clave de la IA.")
        except anthropic.PermissionDeniedError:
            self.log("✖ La clave no tiene permiso o no queda saldo en la cuenta de Anthropic.")
        except anthropic.RateLimitError:
            self.log("✖ Demasiadas peticiones seguidas. Espera un minuto y vuelve a intentarlo.")
        except anthropic.APIConnectionError:
            self.log("✖ No hay conexión a Internet.")
        except anthropic.APIStatusError as e:
            if "credit" in str(e).lower() or "balance" in str(e).lower():
                self.log("✖ No queda saldo en la cuenta de Anthropic.")
            else:
                self.log(f"✖ Error del servicio de IA: {e}")
        except Exception as e:
            self.log(f"✖ Error: {e}")
            traceback.print_exc()
        self.raiz.after(0, self.terminar, creados, coste)

    def terminar(self, creados, coste):
        self.progreso.stop()
        self.btn.config(state="normal")
        if not creados:
            return
        self.log(f"Listo: {len(creados)} Word creado(s). Coste aproximado: {coste:.2f} US$".replace(".", ","))
        if self.abrir_al_final.get():
            for ruta in creados:
                abrir(ruta)
        else:
            messagebox.showinfo("Apuntes a Word", f"¡Listo! {len(creados)} Word creado(s) en:\n{self.carpeta.get()}")


def main():
    args = sys.argv[1:]
    # modo consola para pruebas: apuntes_a_word.py --consola archivo.pdf [carpeta_salida]
    if args and args[0] == "--consola":
        cfg = cargar_config()
        clave = limpiar_clave(os.environ.get("ANTHROPIC_API_KEY") or cfg.get("api_key"))
        opciones = {k: True for k, _ in App.OPCIONES}
        salida = args[2] if len(args) > 2 else carpeta_descargas()
        creados, coste = procesar([args[1]], salida, opciones, clave, print)
        print(f"Coste aproximado: {coste:.2f} US$")
        return
    raiz = TkinterDnD.Tk() if TkinterDnD is not None else tk.Tk()
    try:
        ttk.Style().theme_use("vista")
    except tk.TclError:
        pass
    App(raiz, pdfs_iniciales=args)  # PDFs soltados sobre el icono del programa
    raiz.mainloop()


if __name__ == "__main__":
    main()
